"""PPT Generator — creates PowerPoint decks for SPOC review from Deep Study reports."""
import re
from pathlib import Path
from typing import List, Optional

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUTPUT_DIR = Path("reports") / "pptx"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

DARK_BLUE = RGBColor(0x1A, 0x3C, 0x6E)
ACCENT = RGBColor(0xE8, 0x74, 0x22)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x1A, 0x1A, 0x1A)


def _add_title_slide(prs, module_name: str, lesson_title: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BLUE
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(2))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = module_name
    p.font.size = Pt(28)
    p.font.color.rgb = ACCENT
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = lesson_title
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(12)
    p3 = tf.add_paragraph()
    p3.text = "SPOC Review — Deep Study Report"
    p3.font.size = Pt(18)
    p3.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(24)


def _add_content_slide(prs, title: str, content: str, font_size: int = 16):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), SLIDE_W, Inches(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BLUE
    shape.line.fill.background()
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.7))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = WHITE
    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.333), Inches(5.8))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    lines = content.split('\n')
    first = True
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if first:
            p = tf2.paragraphs[0]
            first = False
        else:
            p = tf2.add_paragraph()
        p.text = line
        p.font.size = Pt(font_size)
        p.font.color.rgb = DARK_TEXT
        p.space_after = Pt(4)


def _add_qa_slide(prs, question: str, answer: str, q_num: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), SLIDE_W, Inches(0.8))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(12.333), Inches(0.6))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = f"POC Question {q_num}"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE
    txBox2 = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(12.333), Inches(1.5))
    tf2 = txBox2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = f"Q: {question}"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = DARK_BLUE
    txBox3 = slide.shapes.add_textbox(Inches(0.5), Inches(2.8), Inches(12.333), Inches(4.3))
    tf3 = txBox3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = f"A: {answer}"
    p3.font.size = Pt(14)
    p3.font.color.rgb = DARK_TEXT


def _extract_section(report: str, section_name: str) -> str:
    """Extract a section from the markdown report."""
    pattern = rf'##\s*{re.escape(section_name)}'
    match = re.search(pattern, report, re.IGNORECASE)
    if not match:
        return ""
    start = match.end()
    next_match = re.search(r'\n##?\s+', report[start:])
    if next_match:
        end = start + next_match.start()
    else:
        end = len(report)
    return report[start:end].strip()


def _extract_poc_questions(report: str) -> list:
    """Extract POC questions and answers from the report.

    Handles several formats the LLM may output:
      1. **Q: What is X?**  /  **A:** ...
      2. **Q1. What is X?**  /  **Ans:** ...
      3. **1. What is X?** followed by an answer paragraph
      4. A numbered list of questions where answers are inline
    """
    section = _extract_section(report, "POC Expected Questions")
    if not section:
        return []

    qa_pairs = []

    # Pattern 1: Explicit Q:/A: blocks (with optional numbering & bold markers)
    #   Q: / Q1. / Q2. / Q1st / Q2nd / Q3rd / Q4th / Q1) ...
    pattern = (
        r'(?:\d+\.?\s*)?\*{1,2}Q(?:\d+)?(?:st|nd|rd|th)?[\.\):]?\s+?(.+?)'
        r'\*{0,2}\s*\n+\s*(?:\d+\.?\s*)?\*{1,2}A(?:nswer|ns|s)?[\.\):]?\s*(.+?)?'
        r'(?=(?:\n(?:\d+\.?\s*)?\*{1,2}Q(?:\d+)?|\Z))'
    )
    matches = re.findall(pattern, section, re.DOTALL | re.IGNORECASE)
    for q, a in matches:
        a = a.strip() if isinstance(a, str) else ""
        q = q.strip()
        if q:
            qa_pairs.append((q, a))

    # Clean leftover markdown bold markers
    cleaned = []
    for q, a in qa_pairs:
        cleaned.append((q.replace("**", "").strip(), a.replace("**", "").strip()))
    qa_pairs = cleaned

    # If no Q:/A: found, try numbered-list-of-questions format.
    # Each item: starts with a digit, question text, then an answer line.
    if not qa_pairs:
        items = re.split(r'\n\s*(?=\d+[\.)]\s)', section)
        for item in items:
            item = item.strip()
            if not item or not re.match(r'^\d+[\.)]\s', item):
                continue  # skip non-numbered chunks (bullets, headers, stray text)
            # First line = question, rest = answer
            lines = [ln.strip() for ln in item.split('\n') if ln.strip()]
            if not lines:
                continue
            q = re.sub(r'^\*+|\*+$', '', lines[0]).strip()
            # Drop leading "Q:" / "Q1:" if present
            q = re.sub(r'^\d+[\.)]\s*', '', q).strip()
            q = re.sub(r'^Q\d*[:.]?\s*', '', q.strip(), flags=re.IGNORECASE)
            a = ' '.join(lines[1:]).strip()
            # A line starting with A:/Ans is the answer
            a = re.sub(r'^(?:A|Ans|Answer)\s*[:.]?\s*', '', a, flags=re.IGNORECASE)
            if q and len(q) > 8:  # avoid tiny fragments
                qa_pairs.append((q, a))

    return qa_pairs


def generate_lesson_ppt(
    module_num: int,
    module_name: str,
    lesson_title: str,
    report_text: str,
    diagram_descriptions: str = "",
    tables: list = None,
) -> str:
    """Generate a PPT for a lesson and save to reports/pptx/.

    Returns the file path of the generated PPT.
    """
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1. Title slide
    _add_title_slide(prs, module_name, lesson_title)

    # 2. Layman's summary
    layman = _extract_section(report_text, "Layman's Explanation")
    if layman:
        _add_content_slide(prs, "Layman's Explanation", layman, font_size=18)

    # 3. Lesson explanation
    explanation = _extract_section(report_text, "Lesson Explanation")
    if not explanation:
        explanation = _extract_section(report_text, "Lab Walkthrough")
    if explanation:
        _add_content_slide(prs, "Key Concepts", explanation[:2000], font_size=14)

    # 4. Diagram descriptions
    if diagram_descriptions:
        _add_content_slide(prs, "Diagrams & Visuals", diagram_descriptions[:2000], font_size=14)

    # 5. Tables
    if tables:
        for i, table in enumerate(tables[:3]):
            _add_content_slide(prs, f"Table {i+1}", table, font_size=12)

    # 6. POC Questions (one per slide)
    qa_pairs = _extract_poc_questions(report_text)
    for i, (q, a) in enumerate(qa_pairs[:10], 1):
        _add_qa_slide(prs, q, a, i)

    # 7. Summary
    summary = _extract_section(report_text, "30-Second Interview Answer")
    if summary:
        _add_content_slide(prs, "30-Second Interview Answer", summary, font_size=18)

    # Save
    safe_name = re.sub(r'[^\w\-]', '_', lesson_title)[:50]
    filename = f"{module_num}_{safe_name}.pptx"
    filepath = OUTPUT_DIR / filename
    prs.save(str(filepath))
    return str(filepath)